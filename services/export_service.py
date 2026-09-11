import os, re
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document as WordDocument
from pptx import Presentation
from pptx.util import Inches

def safe_name(name):
    return re.sub(r"[^a-zA-Z0-9_-]+","_",name)[:80]

def export_pdf(title,content,path):
    styles=getSampleStyleSheet(); story=[]
    for line in content.splitlines():
        if line.startswith("#"): story.append(Paragraph(line.lstrip("# ").replace("&","&amp;"),styles["Heading2"]))
        elif line.strip(): story.append(Paragraph(line.replace("&","&amp;"),styles["BodyText"]))
        story.append(Spacer(1,6))
    SimpleDocTemplate(path,pagesize=A4).build(story)

def export_docx(title,content,path):
    d=WordDocument(); d.add_heading(title,0)
    for line in content.splitlines():
        if line.startswith("## "): d.add_heading(line[3:],2)
        elif line.startswith("# "): d.add_heading(line[2:],1)
        elif line.strip(): d.add_paragraph(line)
    d.save(path)

def export_pptx(title,content,path):
    prs=Presentation(); chunks=[x.strip() for x in content.split("## Slide") if x.strip()]
    if not chunks: chunks=[content]
    for i,chunk in enumerate(chunks):
        slide=prs.slides.add_slide(prs.slide_layouts[1])
        lines=chunk.splitlines()
        slide.shapes.title.text=lines[0].replace(":","").strip() if lines else title
        tf=slide.placeholders[1].text_frame
        tf.text="\n".join(lines[1:])[:1500]
    prs.save(path)
